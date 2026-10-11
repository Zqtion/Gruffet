"""Export everything to one XLSX workbook."""
from __future__ import annotations

import io
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import __version__
from .extract import email_rank
from .images import make_thumbnail
from .store import Store

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
HEAD_FILL = PatternFill("solid", fgColor="1F2A44")
FILLS = {"pass": "C6EFCE", "fail": "FFC7CE", "unknown": "FFEB9C"}
GATE_ORDER = ("website_alive", "identity_match", "not_closed", "phone_corroborated", "email_verified",
              "contact_corroborated")


def _s(v) -> str | int | float | None:
    return _ILLEGAL.sub("", v) if isinstance(v, str) else v


def _ranked(store: Store, lid: str, field: str) -> list[str]:
    """Values ordered by how many independent source families confirm them."""
    vals = store.values(lid, field)
    if field == "email":  # confirmed first, then genuine business inboxes before noisy trust/press/privacy ones
        return sorted(vals, key=lambda v: (-len(store.families_for(lid, field, v)), email_rank(v)))
    return sorted(vals, key=lambda v: -len(store.families_for(lid, field, v)))


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: dict[int, int] | None = None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([_s(v) for v in r])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for i, h in enumerate(headers, 1):
        longest = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, min(max(10, longest + 2), 48))
    return ws


def _link(ws, col: int, url_col_values: bool = True) -> None:
    for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink = c.value
                c.font = Font(color="0563C1", underline="single")


def _signals(store: Store, lid: str, prefix: str) -> str:
    return "; ".join(r["field"].split(":", 1)[1] for r in store.claims(lid) if r["field"].startswith(prefix))


def _socials(store: Store, lid: str) -> str:
    seen = {}
    for c in store.claims(lid):
        if c["field"].startswith("social:"):
            seen.setdefault(c["field"][7:], c["value"])
    return "  ".join(f"{k}: {v}" for k, v in seen.items())


def export(store: Store, path: str | Path, thumbs: bool = True, include_rejected: bool = True) -> dict:
    wb = Workbook()
    wb.remove(wb.active)
    leads = store.leads()
    by_tier = Counter(r["tier"] for r in leads)
    good = [r for r in leads if r["tier"] in ("perfect", "agent-approved")]

    # About
    about = wb.create_sheet("About")
    about_rows = [
        ("gruffet report", f"v{__version__}"), ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("", ""), ("Leads found", len(leads)), ("Perfect (all required gates passed)", by_tier["perfect"]),
        ("Agent-approved (borderline, a human/agent accepted)", by_tier["agent-approved"]),
        ("Awaiting review", by_tier["review"]),
        ("Needs data (found, but no contact details yet; not bad leads)", by_tier["incomplete"]),
        ("Rejected (a check failed)", by_tier["reject"]), ("", ""),
        ("What 'perfect' means",
         "The business website answers, the site is demonstrably this business, no closure signal exists, and a phone "
         "or email is confirmed by two independent sources (e.g. OpenStreetMap AND the business's own site; "
         "email domain has working mail servers)."),
        ("Honest limits",
         "Absence of evidence is not proof. Data comes from OpenStreetMap and public business websites, so coverage "
         "depends on how well a place is mapped. Verify before outreach and follow local marketing/privacy law."),
        ("Images", "Source URLs are recorded for every image. Files belong to their owners; do not republish without permission."),
    ]
    for k, v in about_rows:
        about.append([k, v])
    about.column_dimensions["A"].width = 46
    about.column_dimensions["B"].width = 110
    for r in about.iter_rows():
        r[0].font = Font(bold=True)
        r[1].alignment = Alignment(wrap_text=True, vertical="top")
    about["A1"].font = Font(bold=True, size=14)

    # Perfect Leads
    heads = ["Photo", "Name", "Category", "Phone", "Email", "Website", "Address", "Social", "Tech", "Website signals",
             "Score", "Verified by", "Why", "Lat", "Lon", "ID"]
    rows = []
    for r in good:
        lid = r["id"]
        rows.append(["", r["name"], r["category"], "; ".join(_ranked(store, lid, "phone")[:3]),
                     "; ".join(_ranked(store, lid, "email")[:3]), (store.values(lid, "website") or [""])[0],
                     (_ranked(store, lid, "address") or [""])[0], _socials(store, lid),
                     ", ".join(store.values(lid, "tech")), _signals(store, lid, "signal:"), r["score"],
                     "Gates" if r["tier"] == "perfect" else "Agent", r["tier_reason"], r["lat"], r["lon"], lid])
    ws = _sheet(wb, "Perfect Leads", heads, rows, {1: 14, 2: 30, 4: 28, 5: 32, 6: 34, 7: 40, 8: 50, 9: 34, 10: 40, 13: 40})
    _link(ws, 6)
    if thumbs:
        for i, r in enumerate(good, start=2):
            photo = next((im for im in store.images(r["id"], downloaded_only=True) if im["kind"] != "logo"), None)
            data = make_thumbnail(store.workdir / photo["local_path"]) if photo else None
            if data:
                ws.add_image(XLImage(io.BytesIO(data)), f"A{i}")
                ws.row_dimensions[i].height = 70

    # Review Queue
    rows = []
    for r in store.leads("review"):
        lid = r["id"]
        opens = "; ".join(f"{k}: {v['detail']}" for k, v in store.gates(lid).items()
                          if v["status"] != "pass" and k not in ("phone_corroborated", "email_verified"))
        rows.append([r["name"], r["category"], r["score"], opens, "; ".join(_ranked(store, lid, "phone")[:2]),
                     "; ".join(_ranked(store, lid, "email")[:2]), (store.values(lid, "website") or [""])[0], lid])
    ws = _sheet(wb, "Review Queue", ["Name", "Category", "Score", "Open questions", "Phone", "Email", "Website", "ID"],
                rows, {1: 30, 4: 70})
    _link(ws, 7)

    rows = [[r["name"], r["category"], (store.values(r["id"], "address") or [""])[0],
             "; ".join(store.values(r["id"], "phone")[:2]), (store.values(r["id"], "website") or [""])[0],
             r["tier_reason"], r["lat"], r["lon"], r["id"]] for r in store.leads("incomplete")]
    ws = _sheet(wb, "Needs Data", ["Name", "Category", "Address", "Phone", "Website", "What is missing", "Lat", "Lon", "ID"],
                rows, {1: 30, 3: 40, 6: 50})
    _link(ws, 5)

    if include_rejected:
        _sheet(wb, "Rejected", ["Name", "Category", "Score", "Reason", "ID"],
               [[r["name"], r["category"], r["score"], r["tier_reason"], r["id"]] for r in store.leads("reject")],
               {1: 30, 4: 80})

    # Gate matrix: the cross-check made visible
    ws = _sheet(wb, "Gate Matrix", ["Name", "Tier", *GATE_ORDER, "ID"],
                [[r["name"], r["tier"], *[store.gates(r["id"]).get(g)["status"] if store.gates(r["id"]).get(g) else ""
                                          for g in GATE_ORDER], r["id"]] for r in leads], {1: 30})
    for row_idx, r in enumerate(leads, start=2):
        gates = store.gates(r["id"])
        for col, g in enumerate(GATE_ORDER, start=3):
            cell = ws.cell(row=row_idx, column=col)
            if cell.value in FILLS:
                cell.fill = PatternFill("solid", fgColor=FILLS[cell.value])
                cell.comment = Comment(_s(gates[g]["detail"]) or "", "gruffet")

    # Evidence: every claim with provenance
    ev = []
    for r in leads:
        for c in store.claims(r["id"]):
            if c["field"] != "meta.text_sample":
                ev.append([r["name"], c["field"], c["value"][:300], c["source"], c["label"], r["id"]])
    _sheet(wb, "Evidence", ["Lead", "Field", "Value", "Source", "Label", "ID"], ev, {1: 28, 3: 60, 4: 60})

    # Images
    imgs = []
    names = {r["id"]: r["name"] for r in leads}
    for im in store.images():
        if names.get(im["lead_id"]):
            imgs.append([names[im["lead_id"]], im["kind"], im["url"], im["page_url"], im["local_path"] or "",
                         im["width"], im["height"], im["alt"]])
    ws = _sheet(wb, "Images", ["Lead", "Kind", "Source URL", "Found on page", "Local file", "Width", "Height", "Alt"],
                imgs, {3: 60, 4: 50})
    _link(ws, 3)

    # Competitors (from `gruffet compete`)
    comps = store.competitors()
    if comps:
        crow, cev = [], []
        for c in comps:
            cid = c["id"]
            sig = _signals(store, cid, "site.signal:")
            soc = ", ".join(r["field"].split(":", 1)[1] for r in store.claims(cid) if r["field"].startswith("site.social:"))
            pages = ", ".join(r["field"].split(":", 1)[1] for r in store.claims(cid) if r["field"].startswith("site.page:"))
            crow.append([c["host"], c["status"], "; ".join(store.values(cid, "site.title")[:1]),
                         "; ".join(store.values(cid, "site.description")[:1]), "; ".join(store.values(cid, "site.price")[:6]),
                         "; ".join(store.values(cid, "site.offering")[:6]), ", ".join(store.values(cid, "site.tech")),
                         "; ".join(store.values(cid, "site.content")[:1]), pages, soc, sig,
                         "; ".join(store.values(cid, "site.phone")[:2] + store.values(cid, "site.email")[:2]), c["url"], c["note"]])
            for cl in store.claims(cid):
                cev.append([c["host"], cl["field"], cl["value"][:300], cl["source"], cl["label"]])
        ws = _sheet(wb, "Competitors", ["Competitor", "Status", "Title", "Pitch", "Pricing (heuristic)", "Offers", "Tech stack",
                                        "Content activity", "Key pages", "Social", "Weak spots", "Contact", "URL", "Note"],
                    crow, {1: 24, 3: 34, 4: 50, 5: 40, 6: 44, 7: 36, 8: 50, 11: 40, 13: 36})
        _link(ws, 13)
        _sheet(wb, "Competitor Evidence", ["Competitor", "Field", "Value", "Source", "Label"], cev, {3: 60, 4: 60})

    notes = store.notes()
    if notes:
        ws = _sheet(wb, "Analysis", ["Title", "Analysis (written by the agent from the sourced data)"],
                    [[n["title"], n["body"]] for n in notes], {1: 34, 2: 120})
        for row in ws.iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")

    # Market snapshot (derived from this run's data)
    n = max(len(leads), 1)
    def pct(pred):
        return f"{100 * sum(1 for r in leads if pred(r['id'])) / n:.0f}%"
    snap = [["Businesses found", len(leads)], ["With a website", pct(lambda i: bool(store.values(i, 'website')))],
            ["With a phone", pct(lambda i: bool(store.values(i, 'phone')))],
            ["With an email", pct(lambda i: bool(store.values(i, 'email')))],
            ["With a social page", pct(lambda i: any(c['field'].startswith('social:') for c in store.claims(i)))],
            ["Reachable website", pct(lambda i: (store.gates(i).get('website_alive') or {'status': ''})['status'] == 'pass')],
            ["", ""]]
    for cat, cnt in Counter(r["category"] or "unknown" for r in leads).most_common(15):
        snap.append([f"Category: {cat}", cnt])
    _sheet(wb, "Market Snapshot", ["Metric", "Value"], snap, {1: 40})

    # Run log
    runs = [[datetime.fromtimestamp(r["started_at"]).strftime("%Y-%m-%d %H:%M:%S"), r["command"], r["params"], r["summary"]]
            for r in store.db.execute("SELECT * FROM runs ORDER BY id")]
    _sheet(wb, "Run Log", ["When", "Command", "Parameters", "Summary"], runs, {3: 60, 4: 60})

    wb._sheets.insert(1, wb._sheets.pop(wb._sheets.index(wb["Perfect Leads"])))
    if "Analysis" in wb.sheetnames:
        wb._sheets.insert(1, wb._sheets.pop(wb._sheets.index(wb["Analysis"])))
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    summary = {"file": str(path), "perfect": by_tier["perfect"], "agent_approved": by_tier["agent-approved"],
               "review": by_tier["review"], "incomplete": by_tier["incomplete"], "rejected": by_tier["reject"]}
    store.log_run("export", {"path": str(path)}, summary)
    return summary
