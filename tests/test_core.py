from gruffet import extract as ex
from gruffet.discover import build_query, claims_from_tags
from gruffet.store import Store
from gruffet.util import haversine_m, norm_name, norm_phone, normalize_url, registered_domain
from gruffet.verify import verify_lead


# -- util ------------------------------------------------------------------
def test_norm_name_strips_accents_punct_and_legal_suffix():
    assert norm_name("Café  Zürich, Ltd.") == "cafe zurich"


def test_normalize_url():
    assert normalize_url("Example.com/path#frag") == "https://example.com/path"
    assert normalize_url("mailto:a@b.com") is None
    assert normalize_url("localhost") is None


def test_registered_domain_handles_multi_part_suffix():
    assert registered_domain("https://www.shop.example.co.uk/x") == "example.co.uk"
    assert registered_domain("info.example.com") == "example.com"


def test_norm_phone_e164_and_invalid():
    assert norm_phone("+44 20 7946 0958") == "+442079460958"
    assert norm_phone("0300 1234567", "PK") == "+923001234567"      # national format + region hint
    assert norm_phone("+92 300 1234567") == "+923001234567"         # no hint needed when international
    assert norm_phone("12", "US") is None
    assert norm_phone("0300 1234567") is None                       # national format without a region is ambiguous


def test_haversine():
    assert 100 < haversine_m(0, 0, 0, 0.001) < 120


# -- discover ----------------------------------------------------------------
def test_claims_from_tags():
    tags = {"phone": "+44 20 7946 0958; 123", "contact:email": "Info@Shop.com", "website": "shop.com",
            "addr:housenumber": "5", "addr:street": "High St", "addr:city": "Leeds", "opening_hours": "off",
            "contact:facebook": "shopfb"}
    got = claims_from_tags(tags, "GB")
    assert ("phone", "+442079460958") in got
    assert ("email", "info@shop.com") in got
    assert ("website", "https://shop.com/") in got
    assert ("address", "5 High St, Leeds") in got
    assert ("closed", "opening_hours=off") in got
    assert ("social:facebook", "https://www.facebook.com/shopfb") in got


def test_build_query_known_and_fallback():
    assert '"amenity"="dentist"' in build_query("Dentist", 1.0, 2.0, 500, 50)
    q = build_query("yoga studio", 1.0, 2.0, 500, 50)
    assert '"name"~"yoga.*studio",i' in q and "out center tags 50;" in q


# -- extract ---------------------------------------------------------------
def _cf(email: str, key: int = 0x12) -> str:
    return f"{key:02x}" + "".join(f"{ord(c) ^ key:02x}" for c in email)


HTML = f"""<html><head><title>Smile Dental | Home</title>
<meta property="og:image" content="/img/cover.jpg">
<script type="application/ld+json">{{"@context":"https://schema.org","@graph":[
 {{"@type":"Dentist","name":"Smile Dental","telephone":"+44 20 7946 0958","email":"mailto:hello@smile.example.org",
   "address":{{"@type":"PostalAddress","streetAddress":"5 High St","addressLocality":"Leeds","postalCode":"LS1"}},
   "image":"https://smile.example.org/a.jpg","logo":"https://smile.example.org/logo.png",
   "aggregateRating":{{"ratingValue":"4.8","reviewCount":"120"}}}}]}}</script></head>
<body><h1>Smile Dental</h1>
<a href="mailto:desk@smile.example.org?subject=hi">mail</a>
<span data-cfemail="{_cf('cf@smile.example.org')}">[protected]</span>
<a href="tel:+442079460958">call</a>
<a href="https://www.facebook.com/smiledental">fb</a>
<a href="https://www.facebook.com/sharer/sharer.php?u=x">share</a>
<a href="https://www.linkedin.com/in/some-person">person</a>
<a href="https://www.linkedin.com/company/smile-dental">li</a>
<a href="/contact-us">Contact us</a><a href="/blog">Blog</a>
<img src="/img/team.jpg" width="600" height="400" alt="Our team">
<img src="/img/icon-phone.png" width="600" height="400"><img src="/t.gif" width="1" height="1">
<img data-src="/img/office-logo.png" srcset="/s.jpg 300w, /l.jpg 900w" alt="logo">
<footer>Write to sales [at] smile [dot] example.org</footer></body></html>"""


def test_extract_everything():
    soup = ex.soup_of(HTML)
    text = ex.visible_text(soup)
    biz = ex.business_from_jsonld(ex.jsonld_items(soup))
    assert biz["names"] == ["Smile Dental"] and biz["rating"].startswith("4.8")
    assert "5 High St" in biz["addresses"][0] and biz["emails"] == ["hello@smile.example.org"]

    emails = dict(ex.extract_emails(soup, text))
    assert emails["desk@smile.example.org"] == "EXTRACTED"
    assert emails["cf@smile.example.org"] == "EXTRACTED"          # Cloudflare-obfuscated, decoded
    assert emails["sales@smile.example.org"] == "INFERRED"        # [at]/[dot] de-obfuscation

    phones = dict(ex.extract_phones(soup, text, "GB"))
    assert phones["+442079460958"] == "EXTRACTED"

    soc = ex.extract_socials(soup, "https://smile.example.org/")
    assert soc == {"facebook": "https://facebook.com/smiledental", "linkedin": "https://linkedin.com/company/smile-dental"}

    urls = {i["url"]: i["kind"] for i in ex.extract_images(soup, "https://smile.example.org/")}
    assert urls["https://smile.example.org/img/cover.jpg"] == "cover"
    assert "https://smile.example.org/img/team.jpg" in urls
    assert "https://smile.example.org/img/icon-phone.png" not in urls       # icon filtered
    assert "https://smile.example.org/t.gif" not in urls                     # tracking pixel filtered
    assert urls["https://smile.example.org/l.jpg"] == "logo"                  # largest srcset candidate

    assert ex.find_contact_links(soup, "https://smile.example.org/") == ["https://smile.example.org/contact-us"]


def test_signals():
    assert ex.looks_parked("Buy this domain today")
    assert not ex.looks_parked("Welcome to our clinic")
    assert ex.closed_signals("Sadly we are Permanently Closed since May.")
    assert not ex.closed_signals("Open daily 9-5")


# -- store + verify -----------------------------------------------------------
def make_lead(store, name="Smile Dental", site="https://smile.example.org/"):
    lid, _ = store.upsert_lead(name, "dentist", 53.8, -1.55, "GB")
    store.add_claim(lid, "website", site, "osm:node/1")
    return lid


def test_store_merges_nearby_duplicates(tmp_path):
    s = Store(tmp_path)
    a, new_a = s.upsert_lead("Smile Dental Ltd", "dentist", 53.8000, -1.5500)
    b, new_b = s.upsert_lead("Smile Dental", "dentist", 53.8004, -1.5500)   # ~45 m away, same name
    c, new_c = s.upsert_lead("Smile Dental", "dentist", 53.9000, -1.5500)   # ~11 km away: different branch
    assert new_a and not new_b and a == b and new_c and c != a


def good_site_claims(s, lid):
    src = "site:https://smile.example.org/"
    s.add_claim(lid, "meta.http_status", "200", src)
    s.add_claim(lid, "meta.title", "Smile Dental | Home", src)
    s.add_claim(lid, "meta.text_sample", "Welcome to Smile Dental in Leeds", src)


def test_perfect_lead_needs_independent_corroboration(tmp_path):
    s = Store(tmp_path)
    lid = make_lead(s)
    good_site_claims(s, lid)
    s.add_claim(lid, "phone", "+442079460958", "osm:node/1")
    s.add_claim(lid, "phone", "+442079460958", "site:https://smile.example.org/")
    tier, score = verify_lead(s, lid, {})
    assert tier == "perfect" and score >= 75


def test_single_source_phone_is_only_review(tmp_path):
    s = Store(tmp_path)
    lid = make_lead(s)
    good_site_claims(s, lid)
    s.add_claim(lid, "phone", "+442079460958", "osm:node/1")           # site never confirmed it
    tier, _ = verify_lead(s, lid, {})
    assert tier == "review"


def test_dead_site_rejected_and_agent_cannot_override(tmp_path):
    s = Store(tmp_path)
    lid = make_lead(s)
    s.add_claim(lid, "meta.http_status", "404", "site:https://smile.example.org/")
    s.add_claim(lid, "phone", "+442079460958", "osm:node/1")
    tier, _ = verify_lead(s, lid, {})
    assert tier == "reject"
    import pytest
    with pytest.raises(ValueError):
        from gruffet.review import resolve
        resolve(s, lid, "accept", "looks fine")


def test_email_verified_by_mx_and_domain_match(tmp_path):
    s = Store(tmp_path)
    lid = make_lead(s)
    good_site_claims(s, lid)
    s.add_claim(lid, "email", "hello@smile.example.org", "site:https://smile.example.org/contact")
    assert verify_lead(s, lid, {"smile.example.org": True})[0] == "perfect"
    # A domain with no mail server must NOT corroborate the lead: it falls back to review.
    assert verify_lead(s, lid, {"smile.example.org": False})[0] == "review"
    assert s.gates(lid)["email_verified"]["status"] == "fail"
    # An inconclusive DNS lookup (offline) is "unknown", never a pass.
    assert verify_lead(s, lid, {})[0] == "review"
    assert s.gates(lid)["email_verified"]["status"] == "unknown"


def test_closure_signal_rejects(tmp_path):
    s = Store(tmp_path)
    lid = make_lead(s)
    good_site_claims(s, lid)
    s.add_claim(lid, "phone", "+442079460958", "osm:node/1")
    s.add_claim(lid, "phone", "+442079460958", "site:https://smile.example.org/")
    s.add_claim(lid, "closed", "we are permanently closed", "site:https://smile.example.org/")
    assert verify_lead(s, lid, {})[0] == "reject"


def test_agent_accept_survives_reverify(tmp_path):
    from gruffet.review import resolve
    s = Store(tmp_path)
    lid = make_lead(s)
    good_site_claims(s, lid)
    s.add_claim(lid, "phone", "+442079460958", "osm:node/1")
    assert verify_lead(s, lid, {})[0] == "review"
    assert resolve(s, lid, "accept", "phone checked manually") == "agent-approved"
    assert verify_lead(s, lid, {})[0] == "agent-approved"


def test_export_roundtrip(tmp_path):
    from openpyxl import load_workbook
    from gruffet.export import export
    s = Store(tmp_path)
    lid = make_lead(s)
    good_site_claims(s, lid)
    s.add_claim(lid, "phone", "+442079460958", "osm:node/1")
    s.add_claim(lid, "phone", "+442079460958", "site:https://smile.example.org/")
    verify_lead(s, lid, {})
    out = export(s, tmp_path / "x.xlsx")
    wb = load_workbook(tmp_path / "x.xlsx")
    assert out["perfect"] == 1
    assert {"About", "Perfect Leads", "Review Queue", "Gate Matrix", "Evidence", "Images", "Market Snapshot",
            "Run Log"} <= set(wb.sheetnames)
    assert wb["Perfect Leads"]["B2"].value == "Smile Dental"


def test_overpass_timeout_remark_is_an_error_not_zero_results():
    import pytest
    from gruffet.discover import overpass_elements
    assert overpass_elements({"elements": [{"id": 1}]}) == [{"id": 1}]
    assert overpass_elements({"elements": []}) == []          # genuinely empty is fine
    with pytest.raises(RuntimeError):
        overpass_elements({"elements": [], "remark": "runtime error: Query timed out in \"query\" at line 3"})


def test_no_contact_data_is_incomplete_not_reject(tmp_path):
    s = Store(tmp_path)
    lid, _ = s.upsert_lead("Quiet Dental", "dentist", 53.8, -1.55, "GB")   # OSM knew only the name
    tier, _ = verify_lead(s, lid, {})
    assert tier == "incomplete" and "website" in s.lead(lid)["tier_reason"]


def test_agent_supplied_evidence_goes_through_the_same_gates(tmp_path):
    import pytest
    from gruffet.review import add_fact, missing_cards
    s = Store(tmp_path)
    lid, _ = s.upsert_lead("Quiet Dental", "dentist", 53.8, -1.55, "GB")
    verify_lead(s, lid, {})
    assert lid in missing_cards(s)
    with pytest.raises(ValueError):
        add_fact(s, lid, "phone", "+442079460958")                       # no source URL: refused
    with pytest.raises(ValueError):
        add_fact(s, lid, "website", "not a url")
    add_fact(s, lid, "website", "quietdental.example.org", note="web search")
    add_fact(s, lid, "phone", "020 7946 0958", source_url="https://directory.example.net/quiet-dental")
    assert s.values(lid, "website") == ["https://quietdental.example.org/"]
    assert s.families_for(lid, "phone", "+442079460958") == {"ext"}
    # Site content later confirms the same number -> two independent families -> corroborated.
    s.add_claim(lid, "meta.http_status", "200", "site:https://quietdental.example.org/")
    s.add_claim(lid, "meta.title", "Quiet Dental", "site:https://quietdental.example.org/")
    s.add_claim(lid, "phone", "+442079460958", "site:https://quietdental.example.org/")
    assert verify_lead(s, lid, {})[0] == "perfect"


def test_tiny_page_phone_in_body_is_found_but_long_pages_stay_scoped():
    tiny = ex.soup_of("<html><body><p>Welcome. Call us on 0113 245 4550 today.</p></body></html>")
    assert dict(ex.extract_phones(tiny, ex.visible_text(tiny), "GB")) == {"+441132454550": "INFERRED"}
    filler = "lorem ipsum dolor sit amet " * 200
    long_ = ex.soup_of(f"<html><body><article>{filler} ring 0113 245 4550 </article><footer>no phone</footer></body></html>")
    assert ex.extract_phones(long_, ex.visible_text(long_), "GB") == []     # body noise on long pages is ignored


def test_rename_keeps_old_name_so_rediscovery_does_not_duplicate(tmp_path):
    from gruffet.review import resolve
    s = Store(tmp_path)
    lid = make_lead(s, "Graham Lovis Dental Practice")
    good_site_claims(s, lid)
    s.add_claim(lid, "phone", "+442079460958", "osm:node/1")
    verify_lead(s, lid, {})
    resolve(s, lid, "accept", "same premises and phone; now trades as Paul Calder Dental", rename="Paul Calder Dental")
    assert s.lead(lid)["name"] == "Paul Calder Dental" and s.values(lid, "former_name") == ["Graham Lovis Dental Practice"]
    again, created = s.upsert_lead("Graham Lovis Dental Practice", "dentist", 53.8, -1.55, "GB")  # OSM still has old name
    assert again == lid and not created


def test_email_filtering_and_ranking():
    from gruffet.extract import _email_ok, email_rank
    assert not _email_ok("noreply@shop.com") and not _email_ok("do-not-reply@shop.com")
    assert _email_ok("info@shop.com")
    ranked = sorted(["leedsth-tr.speakup@nhs.net", "j.smith@shop.com", "reception@shop.com"], key=email_rank)
    assert ranked == ["reception@shop.com", "j.smith@shop.com", "leedsth-tr.speakup@nhs.net"]


def test_js_detector_ignores_tiny_static_pages_but_catches_apps():
    static = '<html><head><script>var a=1</script></head><body><h1>Example Domain</h1><p>Short page.</p></body></html>'
    spa = '<html><body><div id="root"></div><script src="/a.js"></script></body></html>'
    nojs = '<html><body><noscript>You need to enable JavaScript to run this app.</noscript></body></html>'
    for html, want in ((static, False), (spa, True), (nojs, True)):
        assert ex.needs_js(html, ex.visible_text(html)) is want
    assert ex.needs_js(spa, "x" * 500) is False          # plenty of real text: never render


def test_js_detector_counts_inline_app_data_but_not_json_ld():
    app = '<html><body><h1>Quotes</h1><script src="/jquery.js"></script><script>var data = [' + '{"a":"b"},' * 100 + ']</script></body></html>'
    ld = '<html><body><h1>Cafe</h1><script type="application/ld+json">{"name":"' + "x" * 800 + '"}</script></body></html>'
    assert ex.needs_js(app, ex.visible_text(app)) is True
    assert ex.needs_js(ld, ex.visible_text(ld)) is False


def test_soup_survives_lone_surrogates():
    assert "ok" in ex.visible_text("<p>ok \udc9d</p>")
