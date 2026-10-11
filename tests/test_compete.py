from datetime import date

from gruffet import extract as ex
from gruffet import techscan as ts
from gruffet.compete import card, find_key_pages, report
from gruffet.store import Store

PRICING = """<html><body><h1>Pricing</h1>
<div class="plan"><h3>Starter</h3><div class="p">$9 <span>/mo</span></div></div>
<div class="plan"><h3>Pro</h3><div class="p">$29 per month</div></div>
<div class="plan"><h3>Enterprise</h3><p>Contact us</p></div>
<div class="plan"><h3>Local</h3><span>PKR 5,000 monthly</span></div></body></html>"""


def test_tech_stack_detection():
    html = ('<link href="/wp-content/themes/x.css"><script src="https://www.googletagmanager.com/gtag/js?id=G-1"></script>'
            '<script src="https://js.stripe.com/v3/"></script><a href="https://wa.me/123">chat</a>')
    assert ts.tech_stack(html) == ["WordPress", "Google Analytics", "Stripe", "WhatsApp chat"]
    assert ts.tech_stack("<html><body>plain</body></html>") == []


def test_site_signals_flag_weak_spots_and_clear_good_sites():
    weak = "<html><body>Hello &copy; 2019 Acme</body></html>"
    sig = ts.site_signals(weak, "http://acme.example.com/", ex.soup_of(weak), today=date(2026, 10, 10))
    assert {"no_https", "no_mobile_viewport", "no_analytics", "no_structured_data", "stale_copyright"} <= set(sig)
    assert "2019" in sig["stale_copyright"]
    good = ('<html><head><meta name="viewport" content="width=device-width">'
            '<script type="application/ld+json">{"@type":"Dentist"}</script>'
            '<script src="https://www.googletagmanager.com/gtag/js?id=G-1"></script>'
            '<script src="https://assets.calendly.com/x.js"></script></head><body>&copy; 2026 Good Ltd</body></html>')
    assert ts.site_signals(good, "https://good.example.com/", ex.soup_of(good), today=date(2026, 10, 10)) == {}


def test_extract_prices_with_plan_names_and_periods():
    prices = {(p["plan"], p["price"], p["period"]) for p in ts.extract_prices(ex.soup_of(PRICING))}
    assert ("Starter", "$9", "/mo") in prices
    assert ("Pro", "$29", "permonth") in prices
    assert any(p[0] == "Local" and p[1].startswith("PKR") for p in prices)
    assert not any(p[0] == "Enterprise" for p in prices)             # "Contact us" has no price


def test_sitemap_parse_and_stats():
    idx = "<sitemapindex><sitemap><loc>https://x.com/s1.xml</loc></sitemap></sitemapindex>"
    assert ts.parse_sitemap(idx) == ([], ["https://x.com/s1.xml"])
    xml = ("<urlset><url><loc>https://x.com/blog/a</loc><lastmod>2026-09-20</lastmod></url>"
           "<url><loc>https://x.com/blog/b</loc><lastmod>2025-01-02T10:00:00Z</lastmod></url>"
           "<url><loc>https://x.com/pricing</loc></url></urlset>")
    urls, children = ts.parse_sitemap(xml)
    st = ts.sitemap_stats(urls, today=date(2026, 10, 10))
    assert children == [] and st["total"] == 3 and st["blog"] == 2
    assert st["newest"] == "2026-09-20" and st["recent_90d"] == 1 and "blog(2)" in st["top_sections"]


def test_find_key_pages_same_site_only():
    html = ('<a href="/pricing">Pricing</a><a href="/pricing/enterprise-plan-details">x</a><a href="/blog/">Blog</a>'
            '<a href="https://other.com/careers">Careers elsewhere</a><a href="/about-us">About</a>')
    pages = find_key_pages(ex.soup_of(html), "https://x.com/")
    assert pages["pricing"] == "https://x.com/pricing"                # shortest path wins
    assert pages["blog"] == "https://x.com/blog/" and pages["about"] == "https://x.com/about-us"
    assert "careers" not in pages                                       # other domain ignored


def test_competitor_workflow_cards_report_and_export(tmp_path):
    from openpyxl import load_workbook
    from gruffet.export import export
    s = Store(tmp_path)
    cid = s.add_competitor("https://www.rival.example.com/", "web search")
    assert cid == "c:rival.example.com" and s.add_competitor("https://rival.example.com/") == cid
    assert len(s.competitors()) == 1
    src = "site:https://rival.example.com/"
    s.add_claim(cid, "site.title", "Rival | Fast leads", src)
    s.add_claim(cid, "site.price", "Pro: $29 permonth", src, "INFERRED")
    s.add_claim(cid, "site.tech", "WordPress", src, "EXTRACTED")
    s.add_claim(cid, "site.signal:no_analytics", "no analytics/pixel tag detected", src, "INFERRED")
    s.add_claim(cid, "site.page:pricing", "https://rival.example.com/pricing", src)
    s.set_competitor_status(cid, "ok")
    text = card(s, s.competitors()[0])
    assert "Pro: $29" in text and "WordPress" in text and "no_analytics" in text and "pricing" in text
    assert "heuristics" in report(s)
    s.add_note("Positioning", "Rival is cheaper but has no analytics tag.")
    export(s, tmp_path / "x.xlsx")
    wb = load_workbook(tmp_path / "x.xlsx")
    assert {"Competitors", "Competitor Evidence", "Analysis"} <= set(wb.sheetnames)
    assert wb.sheetnames[1] == "Analysis"
    assert wb["Competitors"]["A2"].value == "rival.example.com"
    assert "cheaper" in wb["Analysis"]["B2"].value


def test_generic_pricing_heading_is_not_a_plan_name():
    html = '<h1>Pricing</h1><div><span>$35 /mo</span></div><h3>Growth</h3><span>$69 /mo</span>'
    got = {(p["plan"], p["price"]) for p in ts.extract_prices(ex.soup_of(html))}
    assert ("", "$35") in got and ("Growth", "$69") in got


def test_sitemap_freshness_flagged_when_every_url_has_the_same_recent_lastmod():
    urls = [(f"https://x.com/p{i}", "2026-10-09") for i in range(100)]
    assert ts.sitemap_stats(urls, today=date(2026, 10, 10))["lastmod_unreliable"] is True
    mixed = [(f"https://x.com/p{i}", "2024-01-01" if i % 2 else "2026-10-09") for i in range(100)]
    assert ts.sitemap_stats(mixed, today=date(2026, 10, 10))["lastmod_unreliable"] is False


def test_blocked_competitor_card_says_so(tmp_path):
    s = Store(tmp_path)
    cid = s.add_competitor("https://blocked.example.com/")
    s.add_claim(cid, "site.error", "HTTP 403", "site:https://blocked.example.com/")
    s.set_competitor_status(cid, "blocked")
    out = card(s, s.competitors()[0])
    assert "blocked" in out and "refuses automated access" in out


def test_note_with_same_title_replaces_instead_of_duplicating(tmp_path):
    s = Store(tmp_path)
    s.add_note("Comparison", "first draft")
    s.add_note("Comparison", "corrected")
    s.add_note("Other", "x")
    assert [(n["title"], n["body"]) for n in s.notes()] == [("Comparison", "corrected"), ("Other", "x")]
